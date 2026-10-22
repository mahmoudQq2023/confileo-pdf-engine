#!/usr/bin/env python3
"""PDF → PowerPoint (19 Oct 2026). One slide per page, the same size as the page.

Two modes:
  editable  (default) Each slide is the page's pictures and drawings as a background image,
            with the page's text on top as real PowerPoint text boxes — one per paragraph, at
            the paragraph's place, in its size, weight and colour, Arabic right-to-left. The text
            is taken out of the background first (MuPDF redaction of the text only: pictures and
            line art stay), so nothing is printed twice. The words come in reading order from
            pdf_text_rows.page_flow(), the extraction PDF→Word and PDF→Audio use, which turns
            the visual-order Arabic of WeasyPrint/LibreOffice PDFs back into logical order.
  picture   Each slide is the page as one picture: the look is exact, nothing is editable. The
            page's text goes into the slide's notes.

A scanned page (no text, or only an invisible OCR layer) becomes a picture slide in either mode,
with any OCR text in its notes. The presentation takes the first page's size; a page of another
size is fitted inside it, centred. PowerPoint allows 1–56 inches a side, so a larger page is
scaled down and a smaller one up, keeping its shape.

    python3 pdf_to_pptx.py IN.pdf OUT.pptx [--mode editable|picture] [--max-pages 150]

Prints one JSON line. Exit 0 done; 2 unreadable/encrypted; 3 too many pages. Python 3.9 compatible:
PyMuPDF, python-pptx, Pillow.
"""
import argparse
import io
import json
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import warnings

warnings.filterwarnings("ignore", category=DeprecationWarning)
import fitz  # PyMuPDF
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PER_PT = 12700
MIN_SIDE_PT, MAX_SIDE_PT = 72.0, 56 * 72.0
ARABIC = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")
# Faces PowerPoint has (or that the PDF named) with Arabic glyphs; anything else gets Arial for
# the Arabic, which every Windows, Mac and Office Online has.
ARABIC_CAPABLE = {"Arial", "Times New Roman", "Tahoma", "Traditional Arabic", "Simplified Arabic", "Segoe UI",
                  "Sakkal Majalla", "Arabic Typesetting", "Courier New", "Amiri", "Cairo", "Tajawal", "Dubai",
                  "Noto Sans Arabic", "Noto Naskh Arabic", "IBM Plex Sans Arabic", "Scheherazade New"}
FONT_ALIASES = {
    "arialmt": "Arial", "arial": "Arial", "helvetica": "Arial", "timesnewromanpsmt": "Times New Roman",
    "timesnewroman": "Times New Roman", "times": "Times New Roman", "times-roman": "Times New Roman",
    "couriernewpsmt": "Courier New", "courier": "Courier New", "calibri": "Calibri", "cambria": "Cambria",
    "verdana": "Verdana", "tahoma": "Tahoma", "georgia": "Georgia", "segoeui": "Segoe UI",
    "traditionalarabic": "Traditional Arabic", "simplifiedarabic": "Simplified Arabic",
}


def clean_font(name):
    """'ABCDEF+Calibri-BoldItalic' → 'Calibri'; PostScript names to the family PowerPoint knows."""
    n = (name or "").split("+", 1)[-1]
    n = re.split(r"[-,]", n, maxsplit=1)[0]
    n = re.sub(r"(PSMT|MT|PS)$", "", n)
    key = n.replace(" ", "").lower()
    if key in FONT_ALIASES:
        return FONT_ALIASES[key]
    spaced = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", n).strip()
    return spaced or "Arial"


def _texttrace_invisible_share(page):
    """Share of the page's characters drawn invisibly (render mode 3: an OCR layer over a scan)."""
    try:
        spans = page.get_texttrace()
    except Exception:
        return 0.0
    total = invisible = 0
    for s in spans:
        n = len(s.get("chars", ()))
        total += n
        if s.get("type") == 3 or s.get("opacity", 1) == 0:
            invisible += n
    return invisible / total if total else 0.0


def _render(page, dpi, clip=None):
    pix = page.get_pixmap(dpi=dpi, alpha=False, clip=clip)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    small = img.resize((max(1, img.width // 4), max(1, img.height // 4)))
    flat = small.getcolors(maxcolors=6000) is not None
    buf = io.BytesIO()
    if flat:
        img.save(buf, "PNG", optimize=True)
    else:
        img.save(buf, "JPEG", quality=90, subsampling=0, optimize=True)
    buf.seek(0)
    return buf


def _is_bold(font_name, flags):
    f = (font_name or "").lower()
    return bool(flags & 16) or "bold" in f or "black" in f or "heavy" in f


def _runs(row, text):
    """The row's text as runs of one style each (a bold word inside a sentence stays bold).
    Style = (font, flags, colour, size) of the word's first glyph span; a row in one style is one
    run with the row's own text, so its spacing is exactly the extraction's."""
    from pdf_to_word import _normalize_presentation_forms, _tidy
    toks = [t for t in row["tokens"] if t["text"]]

    def style(t):
        sp = t.get("span") or {}
        return (sp.get("font", ""), int(sp.get("flags", 0)), int(sp.get("color", row.get("color", 0))),
                float(t.get("size") or row["size"]))

    styles = [style(t) for t in toks]
    if row.get("rtl") or ARABIC.search(text):
        # One run for a right-to-left line: its style changes are font fallbacks for brackets and
        # digits, and run boundaries move neutral characters (brackets, spaces) under the bidi rules.
        main = max(styles, key=lambda st: sum(1 for x in styles if x == st)) if styles else None
        return [(text, main or ("", 0, row.get("color", 0), row["size"]))]

    def key(st):
        return (clean_font(st[0]), _is_bold(st[0], st[1]), bool(st[1] & 2), st[2], round(st[3] * 2) / 2)
    if not toks or len(set(key(st) for st in styles)) == 1:
        return [(text, styles[0] if styles else ("", 0, row.get("color", 0), row["size"]))]
    runs, cur, cur_key = [], [], None
    for t, st in zip(toks, styles):
        if cur and key(st) != cur_key:
            runs.append((cur, cur_style))
            cur = []
        if not cur:
            cur_key, cur_style = key(st), st
        cur.append(_normalize_presentation_forms(t["text"]))
    runs.append((cur, cur_style))
    out = []
    for i, (words, st) in enumerate(runs):
        piece = _tidy(" ".join(words))
        out.append((piece if i == len(runs) - 1 else piece + " ", st))
    return out


def _text_free_background(doc, pno, dpi):
    """The page rendered with its text removed (pictures and drawings kept)."""
    tmp = fitz.open()
    tmp.insert_pdf(doc, from_page=pno, to_page=pno)
    page = tmp[0]
    any_text = False
    for b in page.get_text("dict").get("blocks", []):
        for ln in b.get("lines", []):
            for sp in ln.get("spans", []):
                if sp.get("text", "").strip():
                    page.add_redact_annot(fitz.Rect(sp["bbox"]), fill=False)
                    any_text = True
    if any_text:
        try:
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
        except (TypeError, AttributeError):
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)
    buf = _render(page, dpi)
    tmp.close()
    return buf


def _segments(row):
    """A row split where the space between two words is wider than 1.25 em — table cells, a label
    and its value, a title and its date — unless the row is justified prose (six words or more,
    every space within a third of the usual one). Each part keeps its words in logical order."""
    toks = [t for t in row["tokens"] if t["text"].strip()]
    if len(toks) < 2:
        return [row]
    size = max(row.get("size") or 0, 1.0)
    visual = sorted(toks, key=lambda t: t["x0"])
    gaps = sorted(max(0.0, b["x0"] - a["x1"]) for a, b in zip(visual, visual[1:]))
    limit = 1.25 * size
    if gaps[-1] <= limit:
        return [row]
    median = gaps[len(gaps) // 2]
    if len(toks) >= 6 and median > 0 and gaps[0] >= 0.66 * median and gaps[-1] <= 1.34 * median:
        return [row]
    groups, cur = [], [visual[0]]
    for a, b in zip(visual, visual[1:]):
        if b["x0"] - a["x1"] > limit:
            groups.append(cur)
            cur = []
        cur.append(b)
    groups.append(cur)
    if len(groups) == 1:
        return [row]
    out = []
    for g in groups:
        ids = set(id(t) for t in g)
        sub = dict(row)
        sub["tokens"] = [t for t in row["tokens"] if id(t) in ids]
        sub["x0"] = min(t["x0"] for t in g)
        sub["x1"] = max(t["x1"] for t in g)
        out.append(sub)
    return out


def _add_text(slide, rows, ox, oy, k, page_rtl, page_w, row_text):
    """One text box per visual paragraph; each PDF line its own paragraph, so nothing reflows.
    A row with wide gaps (a table) becomes one box per cell instead."""
    from pdf_text_rows import paragraphs
    if not rows:
        return 0
    flowing, cells = [], []
    for r in rows:
        parts = _segments(r)
        if len(parts) == 1:
            flowing.append(r)
        else:
            cells.extend([p] for p in parts)
    groups = []
    if flowing:
        groups = paragraphs(flowing, min(r["x0"] for r in flowing), max(r["x1"] for r in flowing))
    groups.extend(cells)
    boxes = 0
    widest = max(r["x1"] - r["x0"] for r in rows)
    for para in groups:
        x0 = min(r["x0"] for r in para)
        x1 = max(r["x1"] for r in para)
        y0 = min(r["y0"] for r in para)
        y1 = max(r["y1"] for r in para)
        size = max(r["size"] for r in para)
        # A little room: PowerPoint's fallback fonts run wider than the PDF's own.
        pad_w = max(4.0, (x1 - x0) * 0.06)
        rtl = sum(1 for r in para if r.get("rtl")) * 2 >= len(para)
        left = x0 - (pad_w if rtl else 0.0)
        tb = slide.shapes.add_textbox(Emu(int((ox + left * k) * EMU_PER_PT)), Emu(int((oy + (y0 - size * 0.12) * k) * EMU_PER_PT)),
                                      Emu(int((x1 - x0 + pad_w) * k * EMU_PER_PT)),
                                      Emu(int(max(y1 - y0 + size * 0.3, size) * k * EMU_PER_PT)))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        pitch = None
        if len(para) > 1:
            gaps = [b["y0"] - a["y0"] for a, b in zip(para, para[1:]) if b["y0"] > a["y0"]]
            if gaps:
                pitch = sorted(gaps)[len(gaps) // 2]
        for i, row in enumerate(para):
            text = row_text(row)
            if not text:
                continue
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            row_rtl = bool(row.get("rtl")) or (page_rtl and ARABIC.search(text) is not None)
            # Centred: a line of a paragraph with equal room on both sides of it; or a lone line
            # on the page's centre (a title), which then grows both ways in a wider fallback font.
            lgap, rgap = row["x0"] - x0, x1 - row["x1"]
            centred = (len(para) > 1 and lgap > 3 and rgap > 3 and abs(lgap - rgap) < 3) or \
                (len(para) == 1 and abs((row["x0"] + row["x1"]) / 2 - page_w / 2) < 3 and (row["x1"] - row["x0"]) < 0.7 * page_w)
            p.alignment = PP_ALIGN.CENTER if centred else (PP_ALIGN.RIGHT if row_rtl else PP_ALIGN.LEFT)
            if row_rtl:
                p._p.get_or_add_pPr().set("rtl", "1")
            if pitch:
                p.line_spacing = Pt(pitch * k)
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            for run_text, style in _runs(row, text):
                run = p.add_run()
                run.text = run_text
                font_name, flags, colour, size = style
                family = clean_font(font_name)
                f = run.font
                f.size = Pt(max(1.0, round(size * k * 2) / 2))
                f.bold = _is_bold(font_name, flags) or (style is None and bool(row.get("bold")))
                f.italic = bool(flags & 2) or "italic" in (font_name or "").lower() or "oblique" in (font_name or "").lower()
                c = int(colour) & 0xFFFFFF
                f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
                f.name = family
                if ARABIC.search(run_text):
                    # PowerPoint draws Arabic with the complex-script face (a:cs), not the Latin one.
                    rpr = run._r.get_or_add_rPr()
                    cs = rpr.find(qn("a:cs"))
                    if cs is None:
                        cs = rpr.makeelement(qn("a:cs"), {})
                        rpr.append(cs)
                    cs.set("typeface", family if family in ARABIC_CAPABLE else "Arial")
                    rpr.set("lang", "ar-EG")
        if centred and len(para) == 1:
            # Widen a centred title symmetrically, so a wider fallback font cannot push it off-centre.
            grow = (x1 - x0) * 0.15
            tb.left = Emu(int((ox + (x0 - grow) * k) * EMU_PER_PT))
            tb.width = Emu(int((x1 - x0 + 2 * grow) * k * EMU_PER_PT))
        boxes += 1
    return boxes


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--mode", choices=("editable", "picture"), default="editable")
    ap.add_argument("--max-pages", type=int, default=150)
    ap.add_argument("--dpi", type=int, default=0)
    a = ap.parse_args(argv)

    def answer(code, **payload):
        print(json.dumps(payload))
        return code

    t0 = time.time()
    try:
        doc = fitz.open(a.input)
    except Exception as e:
        return answer(2, ok=False, reason="unreadable", detail=str(e)[:200])
    if doc.needs_pass:
        return answer(2, ok=False, reason="encrypted")
    if doc.page_count == 0:
        return answer(2, ok=False, reason="unreadable")
    if doc.page_count > a.max_pages:
        return answer(3, ok=False, reason="too_many_pages", pages=doc.page_count, maxPages=a.max_pages)

    from pdf_text_rows import page_flow, row_text_smart

    def row_text(row):
        return row_text_smart(row, chips=False)

    first = doc[0].rect
    scale = 1.0
    if max(first.width, first.height) > MAX_SIDE_PT:
        scale = MAX_SIDE_PT / max(first.width, first.height)
    elif min(first.width, first.height) < MIN_SIDE_PT:
        scale = MIN_SIDE_PT / min(first.width, first.height)
    sw, sh = first.width * scale, first.height * scale
    prs = Presentation()
    prs.slide_width = Emu(int(sw * EMU_PER_PT))
    prs.slide_height = Emu(int(sh * EMU_PER_PT))
    blank = prs.slide_layouts[6]

    editable_dpi = a.dpi or 150
    picture_dpi = a.dpi or 200
    text_boxes = picture_slides = 0
    for pno in range(doc.page_count):
        page = doc[pno]
        r = page.rect
        k = min(sw / r.width, sh / r.height)
        ox, oy = (sw - r.width * k) / 2.0, (sh - r.height * k) / 2.0
        slide = prs.slides.add_slide(blank)

        page_rtl, flow = page_flow(page)
        rows = [row for col_rows, _, _ in flow for row in col_rows]
        scanned = not rows or _texttrace_invisible_share(page) > 0.5
        if a.mode == "picture" or scanned:
            pic = _render(page, picture_dpi)
            slide.shapes.add_picture(pic, Emu(int(ox * EMU_PER_PT)), Emu(int(oy * EMU_PER_PT)),
                                     Emu(int(r.width * k * EMU_PER_PT)), Emu(int(r.height * k * EMU_PER_PT)))
            notes = "\n".join(row_text(row) for row in rows).strip()
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            picture_slides += 1
            continue

        bg = _text_free_background(doc, pno, editable_dpi)
        slide.shapes.add_picture(bg, Emu(int(ox * EMU_PER_PT)), Emu(int(oy * EMU_PER_PT)),
                                 Emu(int(r.width * k * EMU_PER_PT)), Emu(int(r.height * k * EMU_PER_PT)))
        for col_rows, _, _ in flow:
            text_boxes += _add_text(slide, col_rows, ox, oy, k, page_rtl, r.width, row_text)

    prs.save(a.output)
    return answer(0, ok=True, pages=doc.page_count, slides=len(prs.slides), textBoxes=text_boxes,
                  pictureSlides=picture_slides, mode=a.mode, seconds=round(time.time() - t0, 2))


if __name__ == "__main__":
    sys.exit(main())
